#!/usr/bin/env python3
"""
generate_formacion_externa_data.py
Lee FORMACION_EXTERNA.xlsx (Memoria Anual Formación Externa F_FOR_95)
e inyecta los datos en index.html reemplazando el bloque marcado.
"""

import json
import os
import sys
from openpyxl import load_workbook

XLSX_FILE  = 'FORMACION_EXTERNA.xlsx'
INDEX_FILE = 'index.html'
MARKER_START = '/* __FEXT_DATA_START__ */'
MARKER_END   = '/* __FEXT_DATA_END__ */'


def safe_float(v):
    try:
        x = float(v)
        return 0.0 if x != x else x  # NaN check
    except Exception:
        return 0.0


def procesar(ruta):
    print(f'📂  Leyendo {ruta}…')
    wb = load_workbook(ruta, read_only=True)

    # ── EXPEDIENTES TRAMITADOS ──────────────────────────────────────────
    ws = wb['Expedientes tramitados']
    rows = list(ws.iter_rows(values_only=True))
    expedientes = []
    for r in rows[10:16]:
        if r[0] and str(r[0]).strip().lower() != 'total:' and r[1] is not None:
            try:
                expedientes.append({
                    'dir':  str(r[0]).strip(),
                    'pres': int(safe_float(r[1])),
                    'anu':  int(safe_float(r[2])),
                    'den':  int(safe_float(r[3])),
                    'aut':  int(safe_float(r[4])),
                    'desc': int(safe_float(r[5])),
                })
            except Exception:
                pass

    # ── INVERSIÓN ECONÓMICA ─────────────────────────────────────────────
    ws2 = wb['Inversión económica']
    rows2 = list(ws2.iter_rows(values_only=True))
    inversion = []
    for r in rows2[10:17]:
        if r[0] and str(r[0]).strip().lower() != 'total:' and r[1] is not None:
            try:
                ayuda     = safe_float(r[2])
                indirecto = safe_float(r[3])
                total     = safe_float(r[4]) if r[4] and not isinstance(r[4], str) \
                            else ayuda + indirecto
                inversion.append({
                    'dir':       str(r[0]).strip(),
                    'horas':     safe_float(r[1]),
                    'ayuda':     round(ayuda, 2),
                    'indirecto': round(indirecto, 2),
                    'total':     round(total, 2),
                })
            except Exception:
                pass

    # ── TIPO DE PARTICIPACIÓN ───────────────────────────────────────────
    ws3 = wb['Tipo de participaciónn']
    rows3 = list(ws3.iter_rows(values_only=True))
    participacion = []
    for r in rows3[10:17]:
        if r[0] and str(r[0]).strip().lower() != 'total:' and r[1] is not None:
            try:
                participacion.append({
                    'dir':          str(r[0]).strip(),
                    'asistente':    int(safe_float(r[1])),
                    'ponencia':     int(safe_float(r[2])),
                    'comunicacion': int(safe_float(r[3])),
                    'poster':       int(safe_float(r[4])),
                    'otros':        int(safe_float(r[5])),
                    'distancia':    int(safe_float(r[6])),
                })
            except Exception:
                pass

    # ── LUGARES (hojas 14 Dir.*) ────────────────────────────────────────
    dir_map = {
        '14 Dir. Gerencia':     'Gerencia',
        '14 Dir. Médica':       'Médica',
        '14 Dir. Enfermería':   'Enfermería',
        '14 Dir. Gestión_RRHH': 'Gestión/RRHH',
    }
    lugares = []
    for sheet, label in dir_map.items():
        if sheet not in wb.sheetnames:
            continue
        ws4 = wb[sheet]
        for r in list(ws4.iter_rows(values_only=True))[13:]:
            if not r[1] or str(r[1]).strip().lower().startswith('total'):
                continue
            try:
                ayuda     = safe_float(r[3])
                indirecto = safe_float(r[4])
                total     = safe_float(r[5]) if r[5] and not isinstance(r[5], str) \
                            else ayuda + indirecto
                if total > 0:
                    lugares.append({
                        'lugar':     str(r[1]).strip().rstrip(),
                        'dir':       label,
                        'horas':     safe_float(r[2]),
                        'ayuda':     round(ayuda, 2),
                        'indirecto': round(indirecto, 2),
                        'total':     round(total, 2),
                    })
            except Exception:
                pass
    lugares = sorted(lugares, key=lambda x: x['total'], reverse=True)[:20]

    # ── ACTIVIDADES INDIVIDUALES ────────────────────────────────────────
    act_sheets = {
        'lugar-Gerencia':   'Gerencia',
        'lugar-Médica':     'Médica',
        'lugar-Enfermería': 'Enfermería',
        'lugar-Gestión':    'Gestión',
        'lugar-RRHH':       'RRHH',
    }
    actividades = []
    for sheet, label in act_sheets.items():
        if sheet not in wb.sheetnames:
            continue
        ws5 = wb[sheet]
        for r in list(ws5.iter_rows(values_only=True))[12:]:
            if not r[0] or not r[3]:
                continue
            nombre = str(r[0]).strip()
            if nombre.lower().startswith('total'):
                continue
            try:
                horas     = safe_float(r[6])
                ayuda     = safe_float(r[7])
                indirecto = safe_float(r[8])
                total     = safe_float(r[9]) if r[9] and not isinstance(r[9], str) \
                            else ayuda + indirecto
                if str(r[3]).strip() and (horas > 0 or total > 0):
                    actividades.append({
                        'nombre':    nombre,
                        'cat':       str(r[1] or '').strip()[:35],
                        'lugar':     str(r[2] or '').strip()[:40],
                        'actividad': str(r[3]).strip()[:60],
                        'dir':       label,
                        'inicio':    str(r[4] or '').strip()[:12],
                        'horas':     horas,
                        'ayuda':     round(ayuda, 2),
                        'total':     round(total, 2),
                    })
            except Exception:
                pass

    total_pres = sum(e['pres'] for e in expedientes)
    total_aut  = sum(e['aut']  for e in expedientes)
    total_inv  = sum(i['total'] for i in inversion)
    total_h    = sum(i['horas'] for i in inversion)

    print(f'✅  {total_pres} expedientes · {total_aut} autorizados · '
          f'€{total_inv:,.0f} inversión · {total_h:,.0f} horas')

    return {
        'anio':          2025,
        'fuente':        XLSX_FILE,
        'total_pres':    total_pres,
        'total_aut':     total_aut,
        'total_inv':     round(total_inv, 2),
        'total_horas':   total_h,
        'expedientes':   expedientes,
        'inversion':     inversion,
        'participacion': participacion,
        'lugares':       lugares,
        'actividades':   actividades,
    }


def inyectar(data):
    if not os.path.exists(INDEX_FILE):
        print(f'❌  No se encuentra {INDEX_FILE}')
        sys.exit(1)

    with open(INDEX_FILE, 'r', encoding='utf-8') as f:
        html = f.read()

    if MARKER_START not in html or MARKER_END not in html:
        print(f'⚠️   Marcadores no encontrados en {INDEX_FILE}. '
              f'Asegúrate de añadir {MARKER_START} … {MARKER_END} en el código JS.')
        return

    bloque = (
        f'{MARKER_START}\n'
        f'const FEXT_DATA = {json.dumps(data, ensure_ascii=False, separators=(",", ":"))};\n'
        f'{MARKER_END}'
    )

    idx_s = html.index(MARKER_START)
    idx_e = html.index(MARKER_END) + len(MARKER_END)
    html = html[:idx_s] + bloque + html[idx_e:]

    with open(INDEX_FILE, 'w', encoding='utf-8') as f:
        f.write(html)
    print(f'✅  {INDEX_FILE} actualizado con datos de Formación Externa')


if __name__ == '__main__':
    if not os.path.exists(XLSX_FILE):
        print(f'❌  No se encuentra {XLSX_FILE} (recuerda renombrar el archivo)')
        sys.exit(1)
    data = procesar(XLSX_FILE)
    inyectar(data)
    print('🎉  Listo')
